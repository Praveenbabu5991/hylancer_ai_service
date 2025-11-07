# app/utils/file_extractor.py
"""
Utility module for extracting text from various file formats.
Supports: PDF, DOCX, TXT
"""
from typing import BinaryIO
from loguru import logger
import io


async def extract_text_from_file(file_content: bytes, filename: str) -> str:
    """
    Extract text from a file based on its extension.

    Args:
        file_content: The binary content of the file
        filename: The name of the file (used to determine file type)

    Returns:
        Extracted text as a string

    Raises:
        ValueError: If file type is not supported
    """
    file_extension = filename.lower().split('.')[-1]

    if file_extension == 'txt':
        return await extract_text_from_txt(file_content)
    elif file_extension == 'pdf':
        return await extract_text_from_pdf(file_content)
    elif file_extension in ['docx', 'doc']:
        return await extract_text_from_docx(file_content)
    else:
        raise ValueError(
            f"Unsupported file type: .{file_extension}. "
            f"Supported formats: .txt, .pdf, .docx"
        )


async def extract_text_from_txt(file_content: bytes) -> str:
    """
    Extract text from a TXT file.

    Args:
        file_content: The binary content of the text file

    Returns:
        Decoded text content
    """
    try:
        # Try UTF-8 first
        text = file_content.decode('utf-8')
    except UnicodeDecodeError:
        try:
            # Fall back to latin-1
            text = file_content.decode('latin-1')
        except Exception as e:
            logger.error(f"Error decoding text file: {e}")
            raise ValueError(f"Unable to decode text file: {str(e)}")

    return text.strip()


async def extract_text_from_pdf(file_content: bytes) -> str:
    """
    Extract text from a PDF file using PyPDF2.

    Args:
        file_content: The binary content of the PDF file

    Returns:
        Extracted text from all pages
    """
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        raise ImportError(
            "PyPDF2 is required to parse PDF files. "
            "Install it with: pip install PyPDF2"
        )

    try:
        pdf_file = io.BytesIO(file_content)
        pdf_reader = PdfReader(pdf_file)

        text_parts = []
        for page_num, page in enumerate(pdf_reader.pages):
            try:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            except Exception as e:
                logger.warning(f"Error extracting text from page {page_num}: {e}")
                continue

        if not text_parts:
            raise ValueError("No text could be extracted from the PDF file")

        full_text = "\n".join(text_parts)
        return full_text.strip()

    except Exception as e:
        logger.error(f"Error parsing PDF file: {e}")
        raise ValueError(f"Unable to parse PDF file: {str(e)}")


async def extract_text_from_docx(file_content: bytes) -> str:
    """
    Extract text from a DOCX file using python-docx.

    Args:
        file_content: The binary content of the DOCX file

    Returns:
        Extracted text from all paragraphs
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "python-docx is required to parse DOCX files. "
            "Install it with: pip install python-docx"
        )

    try:
        docx_file = io.BytesIO(file_content)
        doc = Document(docx_file)

        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        if not text_parts:
            raise ValueError("No text could be extracted from the DOCX file")

        full_text = "\n".join(text_parts)
        return full_text.strip()

    except Exception as e:
        logger.error(f"Error parsing DOCX file: {e}")
        raise ValueError(f"Unable to parse DOCX file: {str(e)}")
